"""
End-to-end tests for the full spreadsheet pipeline.

Covers: routing, content contract, fallback, WARNING logs, formula-None
threshold unit tests, threshold-triggered fallback, ImportError propagation,
and an optional slow integration test proving data survives the pipeline.

All fixtures are created programmatically with openpyxl -- no binary files
committed to the repository.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

from raganything.parser import MineruParser, _none_ratio_exceeds_threshold


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _create_workbook(tmp_path: Path, sheets: dict[str, list[list[Any]]]) -> Path:
    """
    Build an xlsx file with named sheets from a dict of {sheet_name: 2D cell values}.

    Returns the Path to the saved file.
    """
    wb = openpyxl.Workbook()
    first = True
    for sheet_name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(title=sheet_name)
        for row_idx, row in enumerate(rows, start=1):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
    path = tmp_path / "workbook.xlsx"
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def simple_xlsx(tmp_path: Path) -> Path:
    """Single-sheet workbook with a header row and three data rows."""
    return _create_workbook(
        tmp_path,
        {
            "Sales": [
                ["Region", "Q1", "Q2"],
                ["North", 100, 200],
                ["South", 150, 250],
                ["East", 90, 170],
            ]
        },
    )


@pytest.fixture
def multi_sheet_xlsx(tmp_path: Path) -> Path:
    """Three-sheet workbook, each with distinct data."""
    return _create_workbook(
        tmp_path,
        {
            "Alpha": [["A", "B"], [1, 2]],
            "Beta": [["C", "D"], [3, 4]],
            "Gamma": [["E", "F"], [5, 6]],
        },
    )


# ---------------------------------------------------------------------------
# Test 1 -- direct parser is invoked (parse_office_doc is NOT called)
# ---------------------------------------------------------------------------


class TestDirectParserInvoked:
    def test_direct_parser_invoked(self, simple_xlsx: Path) -> None:
        """parse_spreadsheet() uses SpreadsheetParser directly, not parse_office_doc."""
        parser = MineruParser()

        with patch.object(parser, "parse_office_doc") as mock_office:
            result = parser.parse_spreadsheet(simple_xlsx)

        mock_office.assert_not_called()
        assert isinstance(result, list)
        assert len(result) >= 1

        first = result[0]
        assert "type" in first
        assert "table_body" in first
        assert "table_caption" in first
        assert "page_idx" in first


# ---------------------------------------------------------------------------
# Test 2 -- content contract
# ---------------------------------------------------------------------------


class TestContentContractFields:
    def test_content_contract_fields(self, simple_xlsx: Path) -> None:
        """Each item in the result satisfies the TableModalProcessor contract."""
        result = MineruParser().parse_spreadsheet(simple_xlsx)

        assert len(result) >= 1
        for item in result:
            # Type MUST be "table" -- locked per user decision; not "spreadsheet_table"
            assert item["type"] == "table", (
                f"Expected type 'table', got '{item['type']}'"
            )
            assert isinstance(item["table_body"], str), (
                "table_body must be a str (markdown table)"
            )
            assert "|" in item["table_body"], (
                "table_body must contain pipe characters (GFM table)"
            )
            assert isinstance(item["table_caption"], list), (
                "table_caption must be a list"
            )
            assert len(item["table_caption"]) > 0, (
                "table_caption must contain at least one element"
            )
            assert isinstance(item["page_idx"], int), (
                "page_idx must be an int"
            )


# ---------------------------------------------------------------------------
# Test 3 -- multi-sheet workbook produces multiple items
# ---------------------------------------------------------------------------


class TestMultiSheetProducesMultipleItems:
    def test_multi_sheet_produces_multiple_items(self, multi_sheet_xlsx: Path) -> None:
        """A 3-sheet workbook produces at least 3 table items with correct ordering."""
        result = MineruParser().parse_spreadsheet(multi_sheet_xlsx)

        assert len(result) >= 3, f"Expected >= 3 items, got {len(result)}"

        # page_idx values must be 0, 1, 2 in order (workbook tab order preserved)
        page_indices = [item["page_idx"] for item in result[:3]]
        assert page_indices == [0, 1, 2], (
            f"Expected page_idx [0, 1, 2], got {page_indices}"
        )

        # Each item's caption must include its sheet name
        sheet_names = ["Alpha", "Beta", "Gamma"]
        for item, name in zip(result[:3], sheet_names):
            captions_joined = " ".join(item["table_caption"])
            assert name in captions_joined, (
                f"Sheet name '{name}' not found in table_caption: {item['table_caption']}"
            )


# ---------------------------------------------------------------------------
# Tests 4 + 5 -- fallback on parse failure + WARNING log (combined)
# ---------------------------------------------------------------------------


class TestFallbackOnParseFailure:
    def test_fallback_logs_warning_and_calls_office_doc(
        self, simple_xlsx: Path, caplog: pytest.LogCaptureFixture
    ) -> None:
        """
        When SpreadsheetParser.parse() raises, parse_spreadsheet() falls back to
        parse_office_doc() and logs a WARNING containing the filename.
        """
        sentinel = [{"type": "text", "text": "fallback"}]

        with (
            patch(
                "raganything.spreadsheet.SpreadsheetParser.parse",
                side_effect=ValueError("test error"),
            ),
            patch.object(
                MineruParser,
                "parse_office_doc",
                return_value=sentinel,
            ) as mock_office,
            caplog.at_level(logging.WARNING, logger="raganything.parser"),
        ):
            result = MineruParser().parse_spreadsheet(simple_xlsx)

        # Fallback result is the sentinel
        assert result == sentinel

        # parse_office_doc was called exactly once
        mock_office.assert_called_once()

        # WARNING log must mention the filename
        warning_records = [
            r for r in caplog.records if r.levelno >= logging.WARNING
        ]
        assert len(warning_records) >= 1, "Expected at least one WARNING log record"
        assert any(
            simple_xlsx.name in r.message for r in warning_records
        ), (
            f"Expected WARNING to mention filename '{simple_xlsx.name}'; "
            f"got: {[r.message for r in warning_records]}"
        )
        assert any("falling back" in r.message for r in warning_records), (
            "Expected WARNING to mention 'falling back'"
        )


# ---------------------------------------------------------------------------
# Test 6 -- formula-None threshold unit tests
# ---------------------------------------------------------------------------


class TestFormulaNoneThresholdUnit:
    def _make_table_item(self, rows: list[list[str]]) -> dict:
        """Build a minimal table content_list item from a list of string rows."""
        lines = []
        for i, row in enumerate(rows):
            lines.append("| " + " | ".join(row) + " |")
            if i == 0:
                lines.append("| " + " | ".join(["---"] * len(row)) + " |")
        return {"type": "table", "table_body": "\n".join(lines)}

    def test_threshold_true_when_many_empty_cells(self) -> None:
        """Returns True when >10% of cells are empty."""
        # 9 empty + 1 non-empty = 90% empty -- well above 10%
        content_list = [
            self._make_table_item(
                [
                    ["Header1", "Header2"],
                    ["", ""],
                    ["", ""],
                    ["", ""],
                    ["", ""],
                    ["", "value"],  # only one non-empty among data rows
                ]
            )
        ]
        assert _none_ratio_exceeds_threshold(content_list) is True

    def test_threshold_false_when_few_empty_cells(self) -> None:
        """Returns False when <10% of cells are empty."""
        # 1 empty out of 12 data cells = ~8% -- below 10%
        content_list = [
            self._make_table_item(
                [
                    ["Header1", "Header2", "Header3"],
                    ["a", "b", "c"],
                    ["d", "e", "f"],
                    ["g", "h", ""],  # one empty
                    ["i", "j", "k"],
                ]
            )
        ]
        assert _none_ratio_exceeds_threshold(content_list) is False

    def test_threshold_false_for_empty_content_list(self) -> None:
        """Returns False when content_list is empty (no cells to count)."""
        assert _none_ratio_exceeds_threshold([]) is False

    def test_threshold_ignores_non_table_items(self) -> None:
        """Non-table items do not contribute to the cell count."""
        content_list = [
            {"type": "text", "text": "some text with lots of empty content"},
            self._make_table_item([["A", "B"], ["1", "2"]]),
        ]
        # 0 empty / 4 cells = 0% -- below threshold
        assert _none_ratio_exceeds_threshold(content_list) is False


# ---------------------------------------------------------------------------
# Test 7 -- threshold-triggered fallback in parse_spreadsheet
# ---------------------------------------------------------------------------


class TestThresholdTriggeredFallback:
    def test_formula_none_threshold_triggers_fallback(
        self, tmp_path: Path, caplog: pytest.LogCaptureFixture
    ) -> None:
        """
        When SpreadsheetParser returns >10% empty cells, parse_spreadsheet()
        invokes parse_office_doc() and logs a WARNING about None cells.
        """
        # Build a workbook with mostly empty cells
        path = _create_workbook(
            tmp_path,
            {
                "Formulas": [
                    ["Header1", "Header2"],
                    [None, None],  # empty -- simulates uncached formulas
                    [None, None],
                    [None, None],
                    [None, None],
                    [None, None],
                    ["value", None],  # only one non-empty data cell
                ]
            },
        )
        sentinel = [{"type": "text", "text": "libreoffice_result"}]

        with (
            patch.object(
                MineruParser,
                "parse_office_doc",
                return_value=sentinel,
            ) as mock_office,
            caplog.at_level(logging.WARNING, logger="raganything.parser"),
        ):
            result = MineruParser().parse_spreadsheet(path)

        # parse_office_doc was called because threshold was exceeded
        mock_office.assert_called_once()
        assert result == sentinel

        # WARNING must mention None cells
        warning_messages = [
            r.message for r in caplog.records if r.levelno >= logging.WARNING
        ]
        assert any(
            "None" in msg or "none" in msg.lower() for msg in warning_messages
        ), f"Expected WARNING about None cells, got: {warning_messages}"


# ---------------------------------------------------------------------------
# Test 8 -- ImportError propagates (not swallowed)
# ---------------------------------------------------------------------------


class TestImportErrorPropagates:
    def test_importerror_propagates(
        self, simple_xlsx: Path
    ) -> None:
        """
        ImportError from missing spreadsheet deps is re-raised, not swallowed.

        The import inside parse_spreadsheet() is a local `from raganything.spreadsheet
        import SpreadsheetParser, SpreadsheetConfig`. Setting the module to None in
        sys.modules causes Python to raise ImportError on the from-import.
        """
        import sys

        parser = MineruParser()
        real_module = sys.modules.get("raganything.spreadsheet")

        with patch.object(parser, "parse_office_doc") as mock_office:
            # Temporarily hide the spreadsheet module to simulate missing deps
            sys.modules["raganything.spreadsheet"] = None  # type: ignore[assignment]
            try:
                with pytest.raises(ImportError):
                    parser.parse_spreadsheet(simple_xlsx)
            finally:
                # Restore the real module unconditionally
                if real_module is not None:
                    sys.modules["raganything.spreadsheet"] = real_module
                else:
                    del sys.modules["raganything.spreadsheet"]

        mock_office.assert_not_called()


# ---------------------------------------------------------------------------
# Slow integration test -- data survives the full parsing pipeline
# ---------------------------------------------------------------------------


@pytest.mark.slow
def test_real_lightrag_integration(tmp_path: Path) -> None:
    """
    Verifies that cell values survive the full parsing pipeline end-to-end.

    Marked @pytest.mark.slow so it is skipped in normal CI runs.
    Run explicitly with: pytest -m slow
    """
    path = _create_workbook(
        tmp_path,
        {
            "Finance": [
                ["Metric", "Value"],
                ["Revenue", 42000],
                ["Department", "Engineering"],
                ["Headcount", 15],
            ]
        },
    )

    result = MineruParser().parse_spreadsheet(path)

    assert len(result) >= 1, "Expected at least one content item from parse_spreadsheet"

    # Known cell values must appear somewhere in the table bodies
    all_table_bodies = " ".join(
        item["table_body"] for item in result if item.get("type") == "table"
    )

    assert "42000" in all_table_bodies, (
        f"'42000' not found in table bodies: {all_table_bodies[:500]}"
    )
    assert "Engineering" in all_table_bodies, (
        f"'Engineering' not found in table bodies: {all_table_bodies[:500]}"
    )
