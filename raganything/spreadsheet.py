"""
Spreadsheet processing module for structured tabular data parsing.

Provides helper functions for converting spreadsheet content into
markdown tables suitable for RAG knowledge graph ingestion.
"""

import datetime
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


@dataclass
class SpreadsheetConfig:
    """Configuration for spreadsheet processing."""

    max_rows_per_chunk: int = 150


SUPPORTED_FORMATS: list[str] = [".xlsx", ".xls"]


def get_supported_formats() -> list[str]:
    """Return list of supported spreadsheet file extensions."""
    return SUPPORTED_FORMATS


def is_supported_format(file_path: Path) -> bool:
    """Return True if file extension is a supported spreadsheet format."""
    return Path(file_path).suffix.lower() in SUPPORTED_FORMATS


def sanitize_cell(value: str) -> str:
    """
    Escape markdown special characters in a string cell value.

    Replaces newlines with spaces and escapes pipe characters
    to prevent breaking GFM table structure.
    """
    value = value.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    value = value.replace("|", r"\|")
    return value.strip()


def format_cell_value(value: Any) -> str:
    """
    Coerce an openpyxl cell value to a clean string.

    Handles all Python types returned by openpyxl: None, bool, int,
    float, datetime.datetime, datetime.date, str, and arbitrary types.
    Bool is checked before int because bool is a subclass of int.
    datetime.datetime is checked before datetime.date because datetime
    is a subclass of date.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(value)
    if isinstance(value, datetime.datetime):
        return value.isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, str):
        return sanitize_cell(value)
    return str(value)


def build_merge_map(worksheet: Any) -> dict[tuple[int, int], dict]:
    """
    Build a lookup dict from (row, col) -> merge info for a worksheet.

    Each entry contains:
    - is_anchor: True for the top-left cell of the merged region
    - span_rows: number of rows in the merge
    - span_cols: number of columns in the merge

    Degenerate 1x1 merges (no actual span) are skipped.
    """
    merge_map: dict[tuple[int, int], dict] = {}

    for cell_range in worksheet.merged_cells.ranges:
        span_rows = cell_range.max_row - cell_range.min_row + 1
        span_cols = cell_range.max_col - cell_range.min_col + 1

        if span_rows == 1 and span_cols == 1:
            continue

        for row in range(cell_range.min_row, cell_range.max_row + 1):
            for col in range(cell_range.min_col, cell_range.max_col + 1):
                is_anchor = (row == cell_range.min_row and col == cell_range.min_col)
                merge_map[(row, col)] = {
                    "is_anchor": is_anchor,
                    "span_rows": span_rows,
                    "span_cols": span_cols,
                }

    return merge_map


class SpreadsheetParser:
    """Parse xlsx and xls workbooks into content_list items for RAG ingestion."""

    def __init__(self, config: SpreadsheetConfig | None = None) -> None:
        self.config = config or SpreadsheetConfig()

    def parse(self, file_path: str | Path) -> list[dict]:
        """Parse a spreadsheet file and return content_list items."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")
        ext = path.suffix.lower()
        if ext == ".xlsx":
            return self._parse_xlsx(path)
        if ext == ".xls":
            return self._parse_xls(path)
        raise ValueError(f"Unsupported format '{ext}'. Supported: {SUPPORTED_FORMATS}")

    def _parse_xlsx(self, file_path: Path) -> list[dict]:
        """Parse an xlsx workbook using openpyxl."""
        try:
            import openpyxl
        except ImportError as exc:
            raise ImportError(
                "openpyxl is required for xlsx parsing: pip install openpyxl"
            ) from exc

        wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=False)
        workbook_name = file_path.name
        items: list[dict] = []

        for sheet_idx, ws in enumerate(wb.worksheets):
            if ws.sheet_state != "visible":
                logger.debug("Skipping hidden sheet: %s", ws.title)
                continue

            rows = [[cell.value for cell in row] for row in ws.iter_rows()]

            if all(v is None for row in rows for v in row):
                logger.debug("Skipping empty sheet: %s", ws.title)
                continue

            rows = strip_empty_edges(rows)
            if not rows:
                logger.debug("Skipping empty sheet after edge strip: %s", ws.title)
                continue

            merge_map = build_merge_map(ws)
            items.extend(
                self._chunk_rows(rows, merge_map, ws.title, sheet_idx, workbook_name)
            )

        return items

    def _parse_xls(self, file_path: Path) -> list[dict]:
        """Parse a legacy xls workbook using xlrd."""
        try:
            import xlrd
        except ImportError as exc:
            raise ImportError(
                "xlrd is required for xls parsing: pip install xlrd"
            ) from exc

        book = xlrd.open_workbook(str(file_path))
        workbook_name = file_path.name
        items: list[dict] = []

        for sheet_idx in range(book.nsheets):
            if book.sheet_visibility(sheet_idx) != 0:
                logger.debug("Skipping hidden xls sheet index: %d", sheet_idx)
                continue

            sheet = book.sheet_by_index(sheet_idx)
            rows = self._read_xls_rows(sheet, book)

            if all(v is None for row in rows for v in row):
                logger.debug("Skipping empty xls sheet: %s", sheet.name)
                continue

            rows = strip_empty_edges(rows)
            if not rows:
                continue

            merge_map = self._build_xls_merge_map(sheet)
            items.extend(
                self._chunk_rows(rows, merge_map, sheet.name, sheet_idx, workbook_name)
            )

        return items

    def _read_xls_rows(self, sheet: Any, book: Any) -> list[list[Any]]:
        """Read all rows from an xlrd sheet, resolving date cells."""
        import xlrd

        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    date_tuple = xlrd.xldate_as_tuple(cell.value, book.datemode)
                    row.append(datetime.datetime(*date_tuple) if date_tuple[3:] != (0, 0, 0) else datetime.date(*date_tuple[:3]))
                else:
                    row.append(cell.value if cell.ctype != xlrd.XL_CELL_EMPTY else None)
            rows.append(row)
        return rows

    def _build_xls_merge_map(self, sheet: Any) -> dict[tuple[int, int], dict]:
        """Build merge map from xlrd sheet.merged_cells (0-based, exclusive end -> 1-based inclusive)."""
        merge_map: dict[tuple[int, int], dict] = {}
        for rlo, rhi, clo, chi in sheet.merged_cells:
            span_rows = rhi - rlo
            span_cols = chi - clo
            if span_rows <= 1 and span_cols <= 1:
                continue
            for r in range(rlo, rhi):
                for c in range(clo, chi):
                    is_anchor = (r == rlo and c == clo)
                    merge_map[(r + 1, c + 1)] = {
                        "is_anchor": is_anchor,
                        "span_rows": span_rows,
                        "span_cols": span_cols,
                    }
        return merge_map

    def _render_sheet_to_markdown(
        self, rows: list[list[Any]], merge_map: dict, row_offset: int = 0
    ) -> str:
        """Render a 2D list of cell values into a GFM markdown table string."""
        formatted: list[list[str]] = []
        for row_idx, row in enumerate(rows):
            row_1based = row_offset + row_idx + 1
            formatted_row = []
            for col_idx, value in enumerate(row):
                col_1based = col_idx + 1
                merge_info = merge_map.get((row_1based, col_1based))
                if merge_info is None:
                    formatted_row.append(format_cell_value(value))
                elif not merge_info["is_anchor"]:
                    formatted_row.append("")
                else:
                    cell_str = format_cell_value(value)
                    if merge_info["span_rows"] > 1 or merge_info["span_cols"] > 1:
                        cell_str += f" [merged {merge_info['span_rows']}\u00d7{merge_info['span_cols']}]"
                    formatted_row.append(cell_str)
            formatted.append(formatted_row)

        if not formatted:
            return ""

        col_count = max(len(row) for row in formatted)
        lines = []

        def render_row(row: list[str]) -> str:
            padded = row + [""] * (col_count - len(row))
            return "| " + " | ".join(padded) + " |"

        lines.append(render_row(formatted[0]))
        lines.append("| " + " | ".join(["---"] * col_count) + " |")
        for row in formatted[1:]:
            lines.append(render_row(row))

        return "\n".join(lines)

    def _chunk_rows(
        self,
        rows: list[list[Any]],
        merge_map: dict,
        sheet_name: str,
        sheet_idx: int,
        workbook_name: str,
    ) -> list[dict]:
        """Split rows into chunks and produce content_list items."""
        max_rows = self.config.max_rows_per_chunk
        total = len(rows)

        if total <= max_rows:
            markdown = self._render_sheet_to_markdown(rows, merge_map, row_offset=0)
            caption = f"{workbook_name} - Sheet: {sheet_name}"
            return [{"type": "table", "table_body": markdown, "table_caption": [caption], "page_idx": sheet_idx}]

        items: list[dict] = []
        for start in range(0, total, max_rows):
            chunk = rows[start:start + max_rows]
            end = start + len(chunk)
            markdown = self._render_sheet_to_markdown(chunk, merge_map, row_offset=start)
            caption = f"{workbook_name} - Sheet: {sheet_name} (rows {start + 1}-{end} of {total})"
            items.append({"type": "table", "table_body": markdown, "table_caption": [caption], "page_idx": sheet_idx})

        return items


def strip_empty_edges(rows: list[list[Any]]) -> list[list[Any]]:
    """
    Remove trailing empty columns and rows from a 2D list of cell values.

    Empty means None or "". Only trailing (right/bottom) edges are removed;
    leading rows and columns are preserved to maintain A1 offset semantics.
    Returns [] if the entire grid is empty.
    """
    if not rows:
        return []

    def is_empty_value(v: Any) -> bool:
        return v is None or v == ""

    # Find the last non-empty column index across all rows
    max_col_count = max(len(row) for row in rows)
    last_col = -1
    for col in range(max_col_count):
        for row in rows:
            if col < len(row) and not is_empty_value(row[col]):
                last_col = col
                break

    # Recompute: scan all rows for each col to find rightmost non-empty col
    last_col = -1
    for col in range(max_col_count):
        for row in rows:
            if col < len(row) and not is_empty_value(row[col]):
                last_col = max(last_col, col)

    if last_col == -1:
        return []

    # Trim columns
    trimmed = [row[:last_col + 1] for row in rows]

    # Find the last non-empty row index
    last_row = -1
    for row_idx, row in enumerate(trimmed):
        if not all(is_empty_value(v) for v in row):
            last_row = row_idx

    if last_row == -1:
        return []

    return trimmed[:last_row + 1]
